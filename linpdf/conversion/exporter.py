import os
import csv
from pathlib import Path

import fitz
from PIL import Image

from linpdf.core.document import Document


class Exporter:
    def __init__(self, document: Document):
        self._doc = document

    def _ensure_doc_loaded(self):
        if not self._doc.is_loaded:
            return False, "No document is open", ""
        return None

    def to_text(self, output_path):
        err = self._ensure_doc_loaded()
        if err:
            return err

        src = self._doc._doc
        parts = []
        for i in range(len(src)):
            page = src[i]
            text = page.get_text("text")
            if text.strip():
                parts.append(text)

        output_path = os.path.abspath(output_path)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(parts))

        return True, "Text exported successfully", output_path

    def to_docx(self, output_path):
        err = self._ensure_doc_loaded()
        if err:
            return err

        try:
            from docx import Document as DocxDocument
            from docx.shared import Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            return False, "python-docx is not installed", ""

        docx_doc = DocxDocument()
        src = self._doc._doc

        for i in range(len(src)):
            page = src[i]
            blocks = page.get_text("dict")["blocks"]

            for block in blocks:
                if block["type"] != 0:
                    continue
                for line in block.get("lines", []):
                    para = None
                    for span in line.get("spans", []):
                        text = span.get("text", "").strip()
                        if not text:
                            continue
                        if para is None:
                            para = docx_doc.add_paragraph()
                        run = para.add_run(text)
                        flags = span.get("flags", 0)
                        run.bold = bool(flags & 16)
                        run.italic = bool(flags & 2)
                        size = span.get("size", 11)
                        run.font.size = Pt(size / 72 * 12)
                    if para is not None:
                        bbox = block.get("bbox", [0, 0, 0, 0])
                        page_width = page.rect.width
                        mid = page_width / 2
                        cx = (bbox[0] + bbox[2]) / 2
                        if cx < mid * 0.5:
                            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        elif cx > mid * 1.5:
                            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        else:
                            para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        output_path = os.path.abspath(output_path)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        docx_doc.save(output_path)
        return True, "Document exported to DOCX successfully", output_path

    def to_xlsx(self, output_path):
        err = self._ensure_doc_loaded()
        if err:
            return err

        try:
            import openpyxl
        except ImportError:
            return False, "openpyxl is not installed", ""

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Page 1"
        src = self._doc._doc
        current_sheet = 1

        for i in range(len(src)):
            page = src[i]
            tables = page.find_tables()

            if tables and tables[0].extract():
                for table in tables:
                    for row_data in table.extract():
                        ws.append(list(row_data) if row_data else [])
            else:
                text = page.get_text("text")
                for line in text.split("\n"):
                    ws.append([line])

            if i < len(src) - 1:
                current_sheet += 1
                ws = wb.create_sheet(title=f"Page {current_sheet}")

        output_path = os.path.abspath(output_path)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        wb.save(output_path)
        return True, "Document exported to XLSX successfully", output_path

    def to_pptx(self, output_path):
        err = self._ensure_doc_loaded()
        if err:
            return err

        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
        except ImportError:
            return False, "python-pptx is not installed", ""

        prs = Presentation()
        src = self._doc._doc
        blank_layout = prs.slide_layouts[6]

        for i in range(len(src)):
            page = src[i]
            slide = prs.slides.add_slide(blank_layout)
            text = page.get_text("text")
            if text.strip():
                txBox = slide.shapes.add_textbox(
                    Inches(0.5), Inches(0.5),
                    Inches(9), Inches(7),
                )
                tf = txBox.text_frame
                tf.word_wrap = True
                for j, para_text in enumerate(text.split("\n")[:50]):
                    if j == 0:
                        p = tf.paragraphs[0]
                    else:
                        p = tf.add_paragraph()
                    p.text = para_text

        output_path = os.path.abspath(output_path)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        prs.save(output_path)
        return True, "Document exported to PPTX successfully", output_path

    def to_image(self, output_dir, format="png", dpi=150):
        err = self._ensure_doc_loaded()
        if err:
            return err

        supported = {"png", "jpg", "jpeg", "tif", "tiff"}
        if format not in supported:
            return False, f"Unsupported format '{format}'. Use {supported}", ""

        src = self._doc._doc
        output_dir = os.path.abspath(output_dir)
        os.makedirs(output_dir, exist_ok=True)
        paths = []
        zoom = dpi / 72

        for i in range(len(src)):
            page = src[i]
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
            filename = f"page_{i + 1:04d}.{format}"
            filepath = os.path.join(output_dir, filename)

            if format == "png":
                pix.save(filepath)
            elif format in ("jpg", "jpeg"):
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                img.save(filepath, "JPEG", quality=95)
            elif format in ("tif", "tiff"):
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                img.save(filepath, "TIFF")

            paths.append(filepath)

        return True, f"Exported {len(paths)} page(s) as {format.upper()}", output_dir

    def to_html(self, output_path):
        err = self._ensure_doc_loaded()
        if err:
            return err

        src = self._doc._doc
        parts = [
            "<!DOCTYPE html>",
            "<html>",
            "<head>",
            "<meta charset='utf-8'>",
            "<title>PDF Export</title>",
            "<style>",
            "body{font-family:sans-serif;margin:40px auto;max-width:800px;color:#333;}",
            ".page{margin-bottom:30px;padding:20px;border:1px solid #ddd;border-radius:4px;}",
            ".page h2{font-size:1.2em;color:#666;margin-top:0;}",
            "</style>",
            "</head>",
            "<body>",
        ]

        for i in range(len(src)):
            page = src[i]
            text = page.get_text("text")
            safe = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            html_text = safe.replace("\n", "<br>\n")
            parts.append(f"<div class='page'>")
            parts.append(f"<h2>Page {i + 1}</h2>")
            parts.append(f"<p>{html_text}</p>")
            parts.append(f"</div>")

        parts.extend(["</body>", "</html>"])

        output_path = os.path.abspath(output_path)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(parts))

        return True, "Document exported to HTML successfully", output_path

    def to_csv(self, output_path):
        err = self._ensure_doc_loaded()
        if err:
            return err

        src = self._doc._doc
        output_path = os.path.abspath(output_path)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        with open(output_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Page", "Content"])
            for i in range(len(src)):
                page = src[i]
                text = page.get_text("text")
                for line in text.split("\n"):
                    writer.writerow([i + 1, line])

        return True, "Document exported to CSV successfully", output_path
